"""Repair both B2B workbooks from MAX close reports since a given date.

The operation is deliberately idempotent: affected rows are rebuilt from the
latest MAX message history after timestamped backups are created. It never
sends a message to MAX and never synchronizes jobs to AV Rescue.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
import time
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import openpyxl
try:
    from dotenv import load_dotenv
except ImportError:  # Allows an offline dry-run with an exported message file.
    def load_dotenv():
        return False

load_dotenv()
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from webhook.excel_writer import MONTH_NAMES as REPORT_MONTH_NAMES, append_job
from webhook.payroll_writer import append_salary_row
from webhook.reporting_rules import employee_header, normalize_plate, parse_explicit_closed_report


MOSCOW_TZ = timezone(timedelta(hours=3))
PAGE_SIZE = 100


@dataclass
class RecoveryRecord:
    mid: str
    timestamp_ms: int
    sender: str
    text: str
    job: object


def fetch_messages(oldest_ms: int, newest_ms: int) -> list[dict]:
    import httpx
    from webhook.config import settings

    messages: list[dict] = []
    window_end = newest_ms
    with httpx.Client(timeout=30, verify="/etc/ssl/certs/ca-certificates.crt") as client:
        while True:
            response = client.get(
                f"{settings.MAX_API_BASE}/messages",
                headers={"Authorization": settings.max_bot_token},
                params={"chat_id": settings.max_chat_id, "from": window_end, "to": oldest_ms, "count": PAGE_SIZE},
            )
            response.raise_for_status()
            batch = response.json().get("messages", [])
            if not batch:
                break
            messages.extend(batch)
            if len(batch) < PAGE_SIZE:
                break
            minimum = min(int(item.get("timestamp", 0)) for item in batch)
            if minimum >= window_end or minimum <= oldest_ms:
                break
            window_end = minimum - 1
    by_mid = {}
    for item in messages:
        mid = str(((item.get("body") or {}).get("mid")) or "")
        by_mid[mid or f"without-mid-{len(by_mid)}"] = item
    return sorted(by_mid.values(), key=lambda item: int(item.get("timestamp", 0)))


def message_text(item: dict) -> str:
    body = item.get("body") or {}
    linked_body = ((item.get("link") or {}).get("message") or {})
    return str(body.get("text") or linked_body.get("text") or "")


def sender_name(item: dict) -> str:
    sender = item.get("sender") or {}
    return str(sender.get("first_name") or "").strip()


def correct_region_typo(plate: str, known_plates: list[str]) -> str:
    """Fix only an obvious one-digit region typo, never a plate-letter change."""
    candidates = [
        known for known in known_plates
        if len(plate) == len(known)
        and plate[:6] == known[:6]
        and sum(a != b for a, b in zip(plate[6:], known[6:])) == 1
    ]
    return candidates[0] if len(set(candidates)) == 1 else plate


def build_records(messages: list[dict]) -> tuple[list[RecoveryRecord], list[dict]]:
    known_requests: list[tuple[int, str]] = []
    pending_accepted: dict[str, tuple[int, str, str]] = {}
    records: list[RecoveryRecord] = []
    skipped: list[dict] = []

    for item in messages:
        body = item.get("body") or {}
        mid = str(body.get("mid") or "")
        timestamp_ms = int(item.get("timestamp") or 0)
        text = message_text(item)
        sender = sender_name(item)
        lower = " ".join(text.split()).casefold()
        plate = normalize_plate(text)

        if plate and ("примите" in lower or "новая заявка" in lower):
            known_requests.append((timestamp_ms, plate))

        if "заявка принята" in lower and plate:
            pending_accepted[sender.casefold()] = (timestamp_ms, mid, text)
            continue

        job = parse_explicit_closed_report(text)
        source_mid = mid
        source_text = text
        if job is None and lower.strip() in {"закрыта", "закрылась"}:
            pending = pending_accepted.get(sender.casefold())
            if pending and 0 <= timestamp_ms - pending[0] <= 10 * 60 * 1000:
                source_mid = pending[1] + "+" + mid
                source_text = re.sub(
                    r"заявка\s+принята", "Заявка закрыта", pending[2], count=1, flags=re.IGNORECASE
                )
                job = parse_explicit_closed_report(source_text)

        if job is None:
            if re.search(r"\bзакрыт\w*\b", lower):
                skipped.append({"mid": mid, "reason": "incomplete_close", "text": text[:160]})
            continue

        original_plate = job.license_plate or ""
        recent_plates = [
            request_plate for request_ts, request_plate in known_requests
            if 0 <= timestamp_ms - request_ts <= 3 * 24 * 3600 * 1000
        ]
        job.license_plate = correct_region_typo(original_plate, recent_plates)

        records.append(RecoveryRecord(source_mid, timestamp_ms, sender, source_text, job))

    # Stable message IDs keep the recovery duplicate-safe.
    deduped = {record.mid: record for record in records}
    return sorted(deduped.values(), key=lambda record: record.timestamp_ms), skipped


def workbook_kind(path: Path) -> str:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    for sheet in workbook.worksheets:
        headers = [sheet.cell(1, column).value for column in range(1, min(sheet.max_column, 30) + 1)]
        normalized = [str(value or "").strip().casefold() for value in headers]
        if len(normalized) >= 4 and normalized[:4] == ["дата", "vin/гос.номер тс", "услуга", "сумма"]:
            return "report"
        named_employees = [value for value in normalized[3:] if value and value != "сумма"]
        if normalized[:3] == ["дата", "vin/гос.номер тс", "услуга"] and named_employees:
            return "payroll"
    raise ValueError(f"Cannot identify workbook structure: {path}")


def parse_cell_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%d,%m,%y", "%d,%m,%Y", "%d.%m.%Y", "%d/%m/%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def prune_since(path: Path, sheet_names: set[str], since: date) -> int:
    workbook = openpyxl.load_workbook(path)
    removed = 0
    for sheet in workbook.worksheets:
        if sheet.title.casefold() not in {name.casefold() for name in sheet_names}:
            continue
        rows = [
            row for row in range(2, sheet.max_row + 1)
            if (parse_cell_date(sheet.cell(row, 1).value) or date.min) >= since
        ]
        for row in reversed(rows):
            sheet.delete_rows(row, 1)
            removed += 1
    workbook.save(path)
    return removed


def set_env_value(path: Path, key: str, value: str) -> None:
    lines = path.read_text(encoding="utf-8").splitlines() if path.exists() else []
    replacement = f"{key}={value}"
    found = False
    output = []
    for line in lines:
        if line.startswith(f"{key}="):
            output.append(replacement)
            found = True
        else:
            output.append(line)
    if not found:
        output.append(replacement)
    path.write_text("\n".join(output) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--since", default="2026-09-07")
    parser.add_argument("--messages-json")
    parser.add_argument("--excel-path", default=os.getenv("EXCEL_FILE_PATH", ""))
    parser.add_argument("--payroll-path", default=os.getenv("PAYROLL_FILE_PATH", ""))
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    since_dt = datetime.strptime(args.since, "%Y-%m-%d").replace(tzinfo=MOSCOW_TZ)
    if args.messages_json:
        messages = json.loads(Path(args.messages_json).read_text(encoding="utf-8"))
    else:
        messages = fetch_messages(int(since_dt.timestamp() * 1000), int(time.time() * 1000))
    records, skipped = build_records(messages)

    configured = [Path(args.excel_path), Path(args.payroll_path)]
    kinds = {workbook_kind(path): path for path in configured}
    if set(kinds) != {"report", "payroll"}:
        raise SystemExit("The configured paths must resolve to one report and one payroll workbook")
    report_path = kinds["report"]
    payroll_path = kinds["payroll"]

    preview = {
        "since": args.since,
        "messages": len(messages),
        "records": len(records),
        "skipped_incomplete_closes": len(skipped),
        "report_file": report_path.name,
        "payroll_file": payroll_path.name,
        "rows": [
            {
                "date": datetime.fromtimestamp(record.timestamp_ms / 1000, MOSCOW_TZ).date().isoformat(),
                "plate": record.job.license_plate,
                "service": " + ".join(item.name for item in record.job.services),
                "amount": record.job.total_amount_rub,
                "employee": employee_header(record.sender),
            }
            for record in records
        ],
    }
    print(json.dumps(preview, ensure_ascii=False, indent=2))
    if not args.apply:
        return

    backup_dir = Path("data/excel/backups") / datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    backup_dir.mkdir(parents=True, exist_ok=False)
    shutil.copy2(report_path, backup_dir / report_path.name)
    shutil.copy2(payroll_path, backup_dir / payroll_path.name)

    since_date = since_dt.date()
    report_sheet = f"{REPORT_MONTH_NAMES[since_date.month]} {str(since_date.year)[2:]}"
    removed_report = prune_since(report_path, {report_sheet, REPORT_MONTH_NAMES[since_date.month]}, since_date)
    removed_payroll = prune_since(payroll_path, {REPORT_MONTH_NAMES[since_date.month]}, since_date)

    for record in records:
        append_job(report_path, record.job, record.timestamp_ms, record.text)
        append_salary_row(
            payroll_path,
            record.job,
            record.timestamp_ms,
            employee_header(record.sender),
            record.text,
        )

    env_path = Path(".env")
    set_env_value(env_path, "EXCEL_FILE_PATH", str(report_path.resolve()))
    set_env_value(env_path, "PAYROLL_FILE_PATH", str(payroll_path.resolve()))
    marker = Path("data/excel/recovery-2026-09-07.json")
    marker.write_text(
        json.dumps(
            {
                "completed_at_utc": datetime.now(timezone.utc).isoformat(),
                "backup_dir": str(backup_dir),
                "records": len(records),
                "removed_report_rows": removed_report,
                "removed_payroll_rows": removed_payroll,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(json.dumps({"applied": True, "backup_dir": str(backup_dir), "records": len(records)}))


if __name__ == "__main__":
    main()
