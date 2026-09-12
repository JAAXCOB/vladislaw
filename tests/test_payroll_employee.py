from datetime import datetime, timezone

import openpyxl

from webhook.models import Message, User
from webhook.excel_writer import append_job
from webhook.payroll_writer import (
    _match_employee_column,
    append_salary_row,
    ensure_employee_column,
)
from webhook.reporting_rules import employee_header
from webhook.schema import ExtractedJob, ServiceItem


def test_ensure_employee_column_is_idempotent_and_preserves_existing_data(tmp_path):
    payroll_path = tmp_path / "payroll.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Сентябрь"
    sheet.append(["Дата", "VIN/Гос.номер ТС", "Услуга", "Макунин Павел", 750])
    sheet.append(["2026-09-10", "Е523ТТ797", "Эвакуация", 4500, None])
    workbook.save(payroll_path)

    at = datetime(2026, 9, 10, tzinfo=timezone.utc)
    first = ensure_employee_column(payroll_path, "Буревич Антон", at=at)
    second = ensure_employee_column(payroll_path, "Буревич Антон", at=at)

    assert first == ("Сентябрь", 6, True)
    assert second == ("Сентябрь", 6, False)

    saved = openpyxl.load_workbook(payroll_path)
    saved_sheet = saved["Сентябрь"]
    assert saved_sheet.cell(row=1, column=6).value == "Буревич Антон"
    assert saved_sheet.cell(row=2, column=2).value == "Е523ТТ797"
    assert saved_sheet.cell(row=2, column=4).value == 4500
    assert saved_sheet.max_column == 6


def test_full_name_separates_employees_with_same_first_name():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([
        "Дата",
        "VIN/Гос.номер ТС",
        "Услуга",
        "Николай Петров",
        "Николай Большаков",
    ])

    assert _match_employee_column(sheet, "Николай Большаков Водитель Эва") == 5
    assert _match_employee_column(sheet, "Николай") is None


def test_max_sender_keeps_last_name_and_uses_specific_alias():
    message = Message(
        sender=User(user_id=1, first_name="Николай", last_name="Большаков")
    )

    assert message.effective_sender_name() == "Николай Большаков"
    assert employee_header(message.effective_sender_name()) == "Николай Большаков"


def _closed_job(plate: str, amount: int) -> ExtractedJob:
    return ExtractedJob(
        is_closed_job_report=True,
        license_plate=plate,
        services=[ServiceItem(name="Эвакуация", price_rub=amount)],
        total_amount_rub=amount,
        needs_review=False,
    )


def test_edited_message_updates_evacuation_row_without_duplicate(tmp_path):
    report_path = tmp_path / "evacuation.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Сентябрь 26"
    sheet.append(["Дата", "VIN/Гос.номер ТС", "Услуга", "Сумма"])
    workbook.save(report_path)

    timestamp = int(datetime(2026, 9, 12, tzinfo=timezone.utc).timestamp() * 1000)
    append_job(report_path, _closed_job("А123ВС797", 4500), timestamp, message_id="mid-1")
    append_job(
        report_path,
        _closed_job("А123ВС797", 4590),
        timestamp,
        message_id="mid-1",
        is_edited=True,
    )

    saved = openpyxl.load_workbook(report_path)
    saved_sheet = saved["Сентябрь 26"]
    assert saved_sheet.max_row == 2
    assert saved_sheet.cell(row=2, column=4).value == 4590
    assert saved["_MAX_MESSAGE_INDEX"].sheet_state == "veryHidden"


def test_edited_message_updates_payroll_amount_without_duplicate(tmp_path):
    payroll_path = tmp_path / "payroll.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Сентябрь"
    sheet.append(["Дата", "VIN/Гос.номер ТС", "Услуга", "Макунин Павел", 750])
    workbook.save(payroll_path)

    timestamp = int(datetime(2026, 9, 12, tzinfo=timezone.utc).timestamp() * 1000)
    append_salary_row(
        payroll_path,
        _closed_job("А123ВС797", 4500),
        timestamp,
        "Макунин Павел",
        message_id="mid-1",
        is_edited=True,
    )
    append_salary_row(
        payroll_path,
        _closed_job("А123ВС797", 4590),
        timestamp,
        "Макунин Павел",
        message_id="mid-1",
    )

    saved = openpyxl.load_workbook(payroll_path)
    saved_sheet = saved["Сентябрь"]
    assert saved_sheet.max_row == 2
    assert saved_sheet.cell(row=2, column=4).value == 4590
    assert saved["_MAX_MESSAGE_INDEX"].sheet_state == "veryHidden"
