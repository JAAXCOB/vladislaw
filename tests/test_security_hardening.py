import json
import os
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone

import openpyxl
import pytest
from pydantic import ValidationError

os.environ.setdefault("MAX_BOT_TOKEN", "test-token-placeholder")
os.environ.setdefault("MAX_WEBHOOK_SECRET", "test-secret-ABC")
os.environ.setdefault("MAX_CHAT_ID", "987654321")

from webhook import av_rescue_client, event_inbox  # noqa: E402
from webhook.config import Settings  # noqa: E402
from webhook.excel_writer import append_job  # noqa: E402
from webhook.schema import ExtractedJob, ServiceItem  # noqa: E402


def test_settings_fail_closed_for_empty_secret_and_chat() -> None:
    with pytest.raises(ValidationError):
        Settings(max_bot_token="token", max_webhook_secret="", max_chat_id="1")
    with pytest.raises(ValidationError):
        Settings(max_bot_token="token", max_webhook_secret="valid-secret", max_chat_id="")


def test_schema_rejects_contradictory_or_inconsistent_financial_data() -> None:
    with pytest.raises(ValidationError):
        ExtractedJob(is_closed_job_report=True, is_new_job_request=True)
    with pytest.raises(ValidationError):
        ServiceItem(name="Evacuation", price_rub=-1)
    with pytest.raises(ValidationError):
        ExtractedJob(
            is_closed_job_report=True,
            services=[ServiceItem(name="Evacuation", price_rub=100)],
            total_amount_rub=200,
        )


def _report(path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Сентябрь 26"
    sheet.append(["Дата", "VIN/Гос.номер ТС", "Услуга", "Сумма"])
    workbook.save(path)


def test_excel_formula_prefix_is_written_as_text(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    _report(path)
    job = ExtractedJob(
        is_closed_job_report=True,
        license_plate="А123ВС797",
        services=[ServiceItem(name='=HYPERLINK("https://invalid")', price_rub=100)],
        total_amount_rub=100,
    )
    timestamp = int(datetime(2026, 9, 12, tzinfo=timezone.utc).timestamp() * 1000)
    append_job(path, job, timestamp, message_id="formula-test")
    saved = openpyxl.load_workbook(path, data_only=False)
    assert saved["Сентябрь 26"].cell(2, 3).value.startswith("'=")
    assert saved["Сентябрь 26"].cell(2, 3).data_type == "s"


def test_concurrent_excel_writes_are_serialized(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    _report(path)
    timestamp = int(datetime(2026, 9, 12, tzinfo=timezone.utc).timestamp() * 1000)

    def write(index: int) -> None:
        job = ExtractedJob(
            is_closed_job_report=True,
            license_plate=f"TEST{index}",
            services=[ServiceItem(name="Evacuation", price_rub=100 + index)],
            total_amount_rub=100 + index,
        )
        append_job(path, job, timestamp, message_id=f"mid-{index}")

    with ThreadPoolExecutor(max_workers=6) as pool:
        list(pool.map(write, range(12)))
    saved = openpyxl.load_workbook(path)
    assert saved["Сентябрь 26"].max_row == 13


def test_inbox_is_idempotent_and_dead_letters(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(event_inbox.settings, "event_inbox_path", str(tmp_path / "inbox"))
    monkeypatch.setattr(event_inbox.settings, "event_retry_base_seconds", 0)
    monkeypatch.setattr(event_inbox.settings, "event_max_attempts", 2)
    payload = {"update_type": "message_created", "timestamp": 1}
    identifier, created = event_inbox.enqueue(payload)
    assert created is True
    assert event_inbox.enqueue(payload) == (identifier, False)

    def fail(_payload):
        raise RuntimeError("expected")

    assert event_inbox.drain(fail) == 1
    assert event_inbox.drain(fail) == 1
    assert event_inbox.pending_count() == 0
    dead = list((tmp_path / "inbox" / "dead").glob("*.json"))
    assert len(dead) == 1
    assert json.loads(dead[0].read_text(encoding="utf-8"))["attempts"] == 2


def test_partner_queue_is_not_silently_truncated(tmp_path, monkeypatch) -> None:
    path = tmp_path / "partner-queue.json"
    monkeypatch.setattr(av_rescue_client.settings, "av_rescue_sync_queue_path", str(path))
    records = [
        {"payload": {"source_id": str(index), "event": "upsert"}, "attempts": 1, "next_attempt_at": 1}
        for index in range(501)
    ]
    av_rescue_client._write_queue(records)
    assert len(av_rescue_client._read_queue()) == 501
